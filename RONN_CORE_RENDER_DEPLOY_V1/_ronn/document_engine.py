"""RONN's single document extraction interface.

Docling is an optional structure-aware backend beneath this interface. The
existing format-specific parser remains the bounded fallback, so no second
document routing system is introduced.
"""
from __future__ import annotations

import base64
import io
import re
from pathlib import Path

from docling_adapter import extract as docling_extract, status as docling_status, supported as docling_supported

MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
VERSION = "R23-DOCUMENT-ENGINE-2"


def _decode_data_url(data):
    value = str(data or "")
    if "," in value:
        value = value.split(",", 1)[1]
    return base64.b64decode(value)


def _legacy_extract(filename: str, raw: bytes, max_chars: int):
    ext = Path(filename or "").suffix.lower()
    text = ""
    meta = {"filename": filename, "extension": ext, "bytes": len(raw), "backend": "legacy"}

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        parts = []
        for i, page in enumerate(reader.pages[:120]):
            try:
                parts.append(f"\n--- PAGE {i+1} ---\n" + (page.extract_text() or ""))
            except Exception:
                pass
        text = "".join(parts)
        meta["pages"] = len(reader.pages)
    elif ext == ".docx":
        from docx import Document
        document = Document(io.BytesIO(raw))
        text = "\n".join(p.text for p in document.paragraphs)
        meta["paragraphs"] = len(document.paragraphs)
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        chunks = []
        used = 0
        for sheet in workbook.worksheets[:30]:
            header = f"\n--- SHEET: {sheet.title} ---"
            chunks.append(header)
            used += len(header)
            for row in sheet.iter_rows(values_only=True):
                line = "\t".join("" if value is None else str(value) for value in row)
                chunks.append(line)
                used += len(line) + 1
                if used >= max_chars:
                    break
            if used >= max_chars:
                break
        text = "\n".join(chunks)
        meta["sheets"] = workbook.sheetnames
    elif ext == ".pptx":
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(raw))
        chunks = []
        for i, slide in enumerate(presentation.slides):
            chunks.append(f"\n--- SLIDE {i+1} ---")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    chunks.append(shape.text)
        text = "\n".join(chunks)
        meta["slides"] = len(presentation.slides)
    else:
        text = raw.decode("utf-8", errors="replace")

    text = text[:max_chars]
    meta["extracted_chars"] = len(text)
    return {"text": text, "meta": meta}


def extract_document(filename, data_url, max_chars=80000):
    max_chars = max(1000, min(int(max_chars or 80000), 200000))
    raw = _decode_data_url(data_url)
    if len(raw) > MAX_DOCUMENT_BYTES:
        raise ValueError("Document exceeds the 25 MB extraction limit.")

    ext = Path(filename or "").suffix.lower()
    docling_result = None
    if docling_supported(filename):
        docling_result = docling_extract(filename, raw, max_chars=max_chars, max_pages=120)
        if docling_result.get("ok") and str(docling_result.get("text") or "").strip():
            text = str(docling_result.get("text") or "")[:max_chars]
            meta = {
                "filename": filename,
                "extension": ext,
                "bytes": len(raw),
                "backend": "docling",
                "extracted_chars": len(text),
                "fallback_used": False,
            }
            meta.update(docling_result.get("meta") or {})
            return {"text": text, "meta": meta}

    fallback = _legacy_extract(filename, raw, max_chars)
    fallback["meta"]["fallback_used"] = bool(docling_result is not None)
    if docling_result is not None:
        fallback["meta"]["docling_attempt"] = {
            "ok": bool(docling_result.get("ok")),
            "reason": str(docling_result.get("reason") or "")[:120],
        }
    return fallback


def status():
    return {
        "version": VERSION,
        "single_interface": True,
        "max_document_bytes": MAX_DOCUMENT_BYTES,
        "docling": docling_status(),
        "fallback_formats": [".pdf", ".docx", ".xlsx", ".pptx", "text/utf8"],
        "owns_final_answer": False,
    }
